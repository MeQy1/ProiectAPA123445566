from tkinter import filedialog, messagebox
from manager import ScheduleEngine

class ExcelExportMixin:
    def export_excel(self):
        if not self.last_timetable:
            return messagebox.showwarning("Atenție", "Generați mai întâi orarul în Tab-ul 2!")

        filepath = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Fișiere Excel", "*.xlsx")], initialfile="Orar_Universitar.xlsx")
        if not filepath:
            return

        try:
            import openpyxl
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Orar Universitar"
            g_names = [g["name"] for g in self.groups]
            lookup = self.build_cell_lookup()

            headers = ["Ziua", "Perechea", "Săpt."] + g_names
            ws.append(headers)

            c_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col_i in range(1, len(headers) + 1):
                c = ws.cell(row=1, column=col_i)
                c.fill = PatternFill("solid", fgColor="1E3A8A")
                c.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
                c.alignment = c_center
            ws.row_dimensions[1].height = 28

            thin, dashed, thick_b = Side("thin", color="CBD5E1"), Side("dashed", color="94A3B8"), Side("medium", color="1E3A8A")
            fills = {
                "IMPAR": PatternFill("solid", fgColor="FEF3C7"),
                "PAR": PatternFill("solid", fgColor="E0F2FE"),
                "FULL": PatternFill("solid", fgColor="FFFFFF"),
                "DAY": PatternFill("solid", fgColor="E2E8F0")
            }

            curr_r = 2
            for day in ScheduleEngine.DAYS:
                d_start = curr_r
                for pair in ScheduleEngine.PAIRS:
                    rt, rb = curr_r, curr_r + 1
                    ws.row_dimensions[rt].height = ws.row_dimensions[rb].height = 28

                    ws.cell(row=rt, column=1, value=day)
                    p_cell = ws.cell(row=rt, column=2, value=f"Perechea {pair}\n({ScheduleEngine.PAIR_INTERVALS[pair]})")
                    ws.merge_cells(start_row=rt, start_column=2, end_row=rb, end_column=2)
                    p_cell.alignment, p_cell.font = c_center, Font(name="Segoe UI", size=9, bold=True)

                    for r_sub, lbl in ((rt, "Impar"), (rb, "Par")):
                        cs = ws.cell(row=r_sub, column=3, value=lbl)
                        cs.alignment, cs.font = c_center, Font(name="Segoe UI", size=8, italic=True, color="475569")

                    for g_idx, g_name in enumerate(g_names):
                        col_i = 4 + g_idx
                        cd = lookup.get((day, pair, g_name), {})
                        if "FULL" in cd:
                            ct = ws.cell(row=rt, column=col_i, value=cd["FULL"])
                            ws.merge_cells(start_row=rt, start_column=col_i, end_row=rb, end_column=col_i)
                            ct.fill, ct.alignment, ct.font = fills["FULL"], c_center, Font(name="Segoe UI", size=9, color="0F172A")
                        else:
                            for p_key, r_sub in (("IMPAR", rt), ("PAR", rb)):
                                cell = ws.cell(row=r_sub, column=col_i)
                                cell.alignment = c_center
                                if p_key in cd:
                                    cell.value = f"[{p_key}] {cd[p_key]}"
                                    cell.fill = fills[p_key]
                                    cell.font = Font(name="Segoe UI", size=8, bold=True)

                    b_edge = thick_b if pair == 7 else thin
                    for col_i in range(1, len(headers) + 1):
                        ws.cell(row=rt, column=col_i).border = Border(left=thin, right=thin, top=thin, bottom=dashed)
                        ws.cell(row=rb, column=col_i).border = Border(left=thin, right=thin, top=dashed, bottom=b_edge)
                    curr_r += 2

                ws.merge_cells(start_row=d_start, start_column=1, end_row=curr_r - 1, end_column=1)
                dc = ws.cell(row=d_start, column=1)
                dc.fill, dc.font, dc.alignment = fills["DAY"], Font(name="Segoe UI", size=11, bold=True, color="1E3A8A"), c_center

            for col_let, w in zip(("A", "B", "C"), (14, 18, 9)):
                ws.column_dimensions[col_let].width = w
            for g_idx in range(len(g_names)):
                ws.column_dimensions[openpyxl.utils.get_column_letter(4 + g_idx)].width = 34

            wb.save(filepath)
            messagebox.showinfo("Export Excel Finalizat", f"Orarul a fost exportat cu succes în:\n{filepath}")
        except Exception as e:
            messagebox.showerror("Eroare la Export", f"Eroare: {e}")

